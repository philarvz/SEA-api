from __future__ import annotations

from io import BytesIO
from math import ceil
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from apps.academic.models import Subject

from .subject_access import allowed_subject_ids_for_question_user

FIRST_DATA_ROW = 2
LAST_DATA_ROW = 5001

TYPE_OPTIONS_ES = [
    'Selección única',
    'Selección múltiple',
    'Respuesta abierta',
    'Código',
]
DIFF_OPTIONS_ES = ['BAJA', 'MEDIA', 'ALTA']
BLOOM_OPTIONS_ES = [
    'RECORDAR',
    'COMPRENDER',
    'APLICAR',
    'ANALIZAR',
    'EVALUAR',
    'CREAR',
]
WIDTH_INCREASE_FACTOR = 1.4

_BASE_TEMPLATE_WIDTHS = {
    'A': 20,  # materia
    'B': 18,  # tipo
    'C': 42,  # enunciado
    'D': 14,  # dificultad
    'E': 14,  # bloom
    'F': 10,  # puntos
    'G': 28,  # imagen_url
    'H': 18,  # opcion_1
    'I': 18,  # opcion_2
    'J': 18,  # opcion_3
    'K': 18,  # opcion_4
    'L': 12,  # correctas
    'M': 28,  # test_code
    'N': 14,  # language
}

def subjects_for_template_user(user) -> list[Subject]:
    allowed = allowed_subject_ids_for_question_user(user)
    qs = Subject.objects.filter(status=True).order_by('level_number', 'name')
    if allowed is not None:
        qs = qs.filter(pk__in=allowed)
    return list(qs)

def _inline_list_formula(values: list[str]) -> str:
    return '"' + ','.join(values) + '"'

def _add_list_validation(ws, col_letter: str, formula1: str) -> None:
    dv = DataValidation(type='list', formula1=formula1, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{col_letter}{FIRST_DATA_ROW}:{col_letter}{LAST_DATA_ROW}')


def _apply_template_column_widths(ws) -> None:
    for col, width in _BASE_TEMPLATE_WIDTHS.items():
        ws.column_dimensions[col].width = ceil(width * WIDTH_INCREASE_FACTOR)


def _style_header_cell(cell) -> None:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.alignment = Alignment(horizontal='center')
    cell.fill = PatternFill(fill_type='solid', fgColor='1E3A8A')

def _example_data_rows(subjects: list[Subject]) -> list[list]:
    sub = subjects[0].name if subjects else ''
    return [
        [
            sub,
            'Selección única',
            '¿Cuál es la capital de Francia?',
            'MEDIA',
            'RECORDAR',
            1,
            '',
            'París',
            'Londres',
            'Madrid',
            'Roma',
            '1',
            '',
            '',
        ],
        [
            sub,
            'Selección múltiple',
            '¿Cuáles de los siguientes son números primos?',
            'MEDIA',
            'APLICAR',
            1,
            '',
            '2',
            '4',
            '3',
            '9',
            '1,3',
            '',
            '',
        ],
        [
            sub,
            'Respuesta abierta',
            'Explique brevemente qué es una variable en programación.',
            'ALTA',
            'CREAR',
            2,
            '',
            '',
            '',
            '',
            '',
            '',
            '',
            '',
        ],
        [
            sub,
            'Código',
            'Escriba una función suma(a, b) que devuelva la suma de a y b.',
            'MEDIA',
            'APLICAR',
            1,
            '',
            '',
            '',
            '',
            '',
            '',
            'assert suma(2, 3) == 5',
            'python',
        ],
    ]


def build_questions_template_workbook(user) -> bytes:
    subjects = subjects_for_template_user(user)
    wb = Workbook()

    ws_m = wb.active
    ws_m.title = 'Materias'
    ws_m.cell(1, 1, 'materia')
    _style_header_cell(ws_m.cell(1, 1))
    for i, sub in enumerate(subjects, start=2):
        ws_m.cell(row=i, column=1, value=sub.name)
    ws_m.column_dimensions['A'].width = ceil(52 * WIDTH_INCREASE_FACTOR)

    last_sub_row = max(2, 1 + len(subjects))
    sub_name_range = f"='Materias'!$A$2:$A${last_sub_row}"

    wrap = Alignment(wrap_text=True)

    ws = wb.create_sheet('Plantilla')
    headers = [
        'materia',
        'tipo',
        'enunciado',
        'dificultad',
        'bloom',
        'puntos',
        'imagen_url',
        'opcion_1',
        'opcion_2',
        'opcion_3',
        'opcion_4',
        'correctas',
        'test_code',
        'language',
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        _style_header_cell(c)

    for row_offset, row_vals in enumerate(_example_data_rows(subjects), start=2):
        for col, v in enumerate(row_vals, 1):
            ws.cell(row=row_offset, column=col, value=v).alignment = wrap

    _apply_template_column_widths(ws)

    _add_list_validation(ws, 'A', sub_name_range)
    _add_list_validation(ws, 'B', _inline_list_formula(TYPE_OPTIONS_ES))
    _add_list_validation(ws, 'D', _inline_list_formula(DIFF_OPTIONS_ES))
    _add_list_validation(ws, 'E', _inline_list_formula(BLOOM_OPTIONS_ES))

    wb.active = ws

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
