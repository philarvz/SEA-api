import base64
from django.db import transaction
from django.db.models.deletion import ProtectedError
from django.http import HttpResponse
from loguru import logger
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.decorators import action
from rest_framework.parsers import JSONParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.viewsets import ModelViewSet
from drf_spectacular.utils import OpenApiParameter, extend_schema, extend_schema_view
from drf_spectacular.types import OpenApiTypes

from apps.academic.models import Subject
from apps.academic.permissions import IsTeacherOrAdmin

from .models import Question, Answer, CodeQuestion
from .serializers import (
    QuestionSerializer,
    QuestionListSerializer,
)
from .subject_access import allowed_subject_ids_for_question_user
from .excel_upload import (
    normalize_header_row,
    build_column_maps,
    validate_required_columns,
    get_cell,
    parse_difficulty,
    parse_bloom,
    collect_options_from_row,
    resolve_subject_id,
    parse_points,
    worksheet_for_question_import,
    parse_question_type,
    build_error_rows_workbook,
)
from .template_workbook import build_questions_template_workbook
from utils.responses import error_response, success_response


class QuestionPagination:
    page_size = 10
    page_size_query_param = 'page_size'
    max_page_size = 100

    @classmethod
    def paginate(cls, request, queryset):
        from rest_framework.pagination import PageNumberPagination

        class P(PageNumberPagination):
            page_size = cls.page_size
            page_size_query_param = cls.page_size_query_param
            max_page_size = cls.max_page_size

        paginator = P()
        page = paginator.paginate_queryset(queryset, request)
        return page, paginator


def _paginated_payload(request, queryset, serializer_class):
    page, paginator = QuestionPagination.paginate(request, queryset)
    if page is None:
        ser = serializer_class(queryset, many=True)
        count = queryset.count()
        return {
            'results': ser.data,
            'pagination': {
                'count': count,
                'page': 1,
                'page_size': count,
                'total_pages': 1,
                'next': None,
                'previous': None,
            },
        }
    ser = serializer_class(page, many=True)
    page_size = paginator.get_page_size(request) or paginator.page_size
    return {
        'results': ser.data,
        'pagination': {
            'count': paginator.page.paginator.count,
            'page': paginator.page.number,
            'page_size': page_size,
            'total_pages': paginator.page.paginator.num_pages,
            'next': paginator.get_next_link(),
            'previous': paginator.get_previous_link(),
        },
    }


def _validate_upload_file(up):
    if not up:
        return 'Debe adjuntar un archivo .xlsx en el campo "file".'
    if not up.name.lower().endswith('.xlsx'):
        return 'Solo se aceptan archivos .xlsx.'
    # Limit file size to 10 MB to prevent zip-bomb / DoS
    max_size = 10 * 1024 * 1024
    if up.size > max_size:
        return 'El archivo excede el tamaño máximo permitido (10 MB).'
    return None


def _load_excel_rows(up):
    try:
        wb = load_workbook(up, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning('Excel upload parse error: {}', exc)
        return None, 'No se pudo leer el archivo Excel.'
    ws = worksheet_for_question_import(wb)
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        return None, 'El archivo está vacío.'
    return rows, None


def _resolve_correct_indices(options, correct_raw):
    correct_indices = set()
    for part in correct_raw.split(','):
        part = part.strip()
        if not part:
            continue
        if part.isdigit():
            correct_indices.add(int(part) - 1)
        else:
            for i, opt in enumerate(options):
                if opt.lower() == part.lower():
                    correct_indices.add(i)
                    break
    return correct_indices


def _create_choice_answers(q, qtype, options, correct_raw):
    if len(options) < 2 or len(options) > 4:
        raise ValueError('Se requieren entre 2 y 4 opciones (separadas por coma).')
    correct_indices = _resolve_correct_indices(options, correct_raw)
    if qtype == 'MULTIPLE_CHOICE' and len(correct_indices) != 1:
        raise ValueError('Para el tipo de pregunta de selección única indique exactamente un índice correcto (1..n).')
    if qtype == 'MULTIPLE_SELECTION' and len(correct_indices) < 1:
        raise ValueError('Para el tipo de pregunta de selección múltiple indique al menos una respuesta correcta.')
    for i, opt in enumerate(options):
        Answer.objects.create(
            id_question=q,
            answer_text=opt,
            is_correct=i in correct_indices,
        )


def _create_code_question_obj(q, test_code, language):
    tests = [test_code] if test_code else []
    if not tests:
        raise ValueError('test_code es obligatorio para el tipo de pregunta de código.')
    if not language:
        raise ValueError('language es obligatorio para el tipo de pregunta de código.')
    CodeQuestion.objects.create(
        question=q,
        language=language,
        test_cases=tests,
    )


def _create_type_specific_objects(q, qtype, options, correct_raw, test_code, language):
    if qtype in ('MULTIPLE_CHOICE', 'MULTIPLE_SELECTION'):
        _create_choice_answers(q, qtype, options, correct_raw)
    elif qtype == 'CODE':
        if options:
            raise ValueError('Las preguntas de tipo código no deben tener opciones.')
        _create_code_question_obj(q, test_code, language)
    elif qtype == 'OPEN':
        if options:
            raise ValueError('Las preguntas de tipo abierto no deben tener opciones.')
    else:
        raise ValueError(f'Tipo no soportado: {qtype}')


def _process_upload_row(row, col_by_canon, col_by_header, user):
    """Parse one Excel row, create and return the Question with related objects."""
    text = get_cell(row, col_by_canon.get('question_text'))
    if not text:
        raise ValueError('El enunciado (question_text / enunciado) está vacío.')

    qtype = parse_question_type(get_cell(row, col_by_canon.get('type'), 'MULTIPLE_CHOICE'))
    subject_id = resolve_subject_id(row, col_by_canon, user)

    options = collect_options_from_row(row, col_by_header)
    if not options and 'options' in col_by_canon:
        options_raw = get_cell(row, col_by_canon['options'], '') or ''
        options = [p.strip() for p in str(options_raw).split(',') if p.strip()]

    correct_raw = str(get_cell(row, col_by_canon.get('correct_answers'), '') or '')
    difficulty = parse_difficulty(get_cell(row, col_by_canon.get('difficulty'), 'medium'))
    bloom = parse_bloom(get_cell(row, col_by_canon.get('bloom_level'), 'remember'))
    image_url = get_cell(row, col_by_canon.get('image_url'), '') or ''
    test_code = get_cell(row, col_by_canon.get('test_code'), '') or ''
    language = str(get_cell(row, col_by_canon.get('language'), '') or '').strip().lower()
    points = parse_points(get_cell(row, col_by_canon.get('points'), 1))

    with transaction.atomic():
        q = Question.objects.create(
            id_subject_id=subject_id,
            statement=str(text),
            question_type=qtype,
            difficulty=difficulty,
            bloom_level=bloom,
            image_url=image_url or None,
            points=points,
            status=True,
        )
        _create_type_specific_objects(q, qtype, options, correct_raw, test_code, language)
    return q


def _is_empty_excel_row(row):
    return not row or all(v is None or str(v).strip() == '' for v in row)


def _normalize_subject_label(value):
    return str(value or '').strip()


def _subject_name_from_row(row, col_by_canon):
    subject_name_idx = col_by_canon.get('subject_name')
    if subject_name_idx is not None:
        subject_name = _normalize_subject_label(get_cell(row, subject_name_idx, ''))
        if subject_name:
            return subject_name

    subject_id_idx = col_by_canon.get('subject_id')
    if subject_id_idx is None:
        return ''

    raw_id = get_cell(row, subject_id_idx, '')
    try:
        subject_id = int(raw_id)
    except (TypeError, ValueError):
        return ''

    subject = Subject.objects.filter(pk=subject_id, status=True).first()
    return subject.name if subject else ''


def _add_if_present(values: set[str], value: str) -> None:
    if value:
        values.add(value)


def _process_subject_row_for_admin(row, col_by_canon, valid_subject_names: set[str]) -> None:
    _add_if_present(valid_subject_names, _subject_name_from_row(row, col_by_canon))


def _resolve_subject_for_teacher(row, col_by_canon):
    try:
        return resolve_subject_id(row, col_by_canon, None)
    except Exception:
        return None


def _process_subject_row_for_teacher(
    row,
    col_by_canon,
    allowed_ids,
    valid_subject_names: set[str],
    denied_subject_names: set[str],
) -> None:
    sid = _resolve_subject_for_teacher(row, col_by_canon)
    if sid is None:
        return
    name = _subject_name_from_row(row, col_by_canon)
    if sid in allowed_ids:
        _add_if_present(valid_subject_names, name)
        return
    _add_if_present(denied_subject_names, name)


def _subject_access_metadata(rows, col_by_canon, allowed_ids):
    valid_subject_names = set()
    denied_subject_names = set()
    is_admin_scope = allowed_ids is None
    for row in rows[1:]:
        if _is_empty_excel_row(row):
            continue
        if is_admin_scope:
            _process_subject_row_for_admin(row, col_by_canon, valid_subject_names)
            continue
        _process_subject_row_for_teacher(
            row,
            col_by_canon,
            allowed_ids,
            valid_subject_names,
            denied_subject_names,
        )
    partial = bool(valid_subject_names and denied_subject_names)
    no_access = bool(not valid_subject_names and denied_subject_names)
    return {
        'valid_subjects': sorted(valid_subject_names),
        'denied_subjects': sorted(denied_subject_names),
        'partial_subject_access': partial,
        'no_subject_access': no_access,
    }


_PATH_ID = OpenApiParameter('id', OpenApiTypes.INT, OpenApiParameter.PATH, description='ID de la pregunta')


@extend_schema_view(
    retrieve=extend_schema(parameters=[_PATH_ID]),
    update=extend_schema(parameters=[_PATH_ID]),
    partial_update=extend_schema(parameters=[_PATH_ID]),
    destroy=extend_schema(parameters=[_PATH_ID]),
)
class QuestionViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated, IsTeacherOrAdmin]
    parser_classes = [JSONParser, MultiPartParser]
    lookup_field = 'id_question'
    lookup_url_kwarg = 'pk'

    def get_queryset(self):
        qs = (
            Question.objects.select_related('id_subject', 'modified_by')
            .prefetch_related('answers')
            .order_by('-id_question')
        )
        allowed = allowed_subject_ids_for_question_user(self.request.user)
        if allowed is not None:
            qs = qs.filter(id_subject_id__in=allowed)
        qtype = self.request.query_params.get('type')
        diff = self.request.query_params.get('difficulty')
        subject = self.request.query_params.get('id_subject')
        if qtype:
            qs = qs.filter(question_type=qtype)
        if diff:
            qs = qs.filter(difficulty=diff)
        if subject:
            qs = qs.filter(id_subject_id=subject)
        return qs

    def get_serializer_class(self):
        if self.action == 'list':
            return QuestionListSerializer
        return QuestionSerializer

    def list(self, request, *args, **kwargs):
        queryset = self.filter_queryset(self.get_queryset())
        payload = _paginated_payload(request, queryset, QuestionListSerializer)
        return success_response(payload)

    def retrieve(self, request, *args, **kwargs):
        instance = self.get_object()
        ser = QuestionSerializer(instance)
        return success_response(ser.data)

    def create(self, request, *args, **kwargs):
        ser = QuestionSerializer(data=request.data, context={'request': request})
        ser.is_valid(raise_exception=True)
        instance = ser.save()
        logger.info('Question created | user={} id={}', request.user.pk, instance.pk)
        out = QuestionSerializer(instance)
        return success_response(out.data, 'Pregunta creada.', status_code=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        ser = QuestionSerializer(instance, data=request.data, partial=partial, context={'request': request})
        ser.is_valid(raise_exception=True)
        instance = ser.save()
        logger.info('Question updated | user={} id={}', request.user.pk, instance.pk)
        return success_response(QuestionSerializer(instance).data, 'Pregunta actualizada.')

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        try:
            instance.delete()
        except ProtectedError:
            return error_response(
                'La pregunta está asociada a un examen y no puede eliminarse.',
                status_code=status.HTTP_400_BAD_REQUEST,
            )
        logger.info('Question deleted | user={} id={}', request.user.pk, kwargs.get('pk'))
        return success_response({'id': instance.pk}, 'Pregunta eliminada.')

    @extend_schema(
        summary='Subir preguntas (Excel)',
        parameters=[
            OpenApiParameter('file', OpenApiTypes.BINARY, required=True),
        ],
    )
    @action(detail=False, methods=['post'], url_path='upload', parser_classes=[MultiPartParser])
    def upload(self, request):
        up = request.FILES.get('file')
        file_err = _validate_upload_file(up)
        if file_err:
            return error_response(file_err)

        rows, load_err = _load_excel_rows(up)
        if load_err:
            return error_response(load_err)

        headers_norm = normalize_header_row(rows[0])
        col_by_canon, col_by_header = build_column_maps(headers_norm)
        req_err = validate_required_columns(col_by_canon)
        if req_err:
            return error_response('El formato del archivo no es correcto')

        access_info = _subject_access_metadata(rows, col_by_canon, allowed_subject_ids_for_question_user(request.user))
        if access_info['no_subject_access']:
            return error_response('No tienes acceso a las materias del archivo.')

        total = 0
        created = 0
        errors = []
        rows_with_errors = []

        for row_idx, row in enumerate(rows[1:], start=2):
            if _is_empty_excel_row(row):
                continue
            total += 1
            try:
                _process_upload_row(row, col_by_canon, col_by_header, request.user)
                created += 1
            except Exception as exc:
                row_error = {'row': row_idx, 'error': str(exc)}
                errors.append(row_error)
                rows_with_errors.append({
                    'row_number': row_idx,
                    'row_data': row,
                    'error': row_error['error'],
                })

        errors_file = None
        if rows_with_errors:
            wb_data = build_error_rows_workbook(rows[0], rows_with_errors)
            errors_file = {
                'filename': 'preguntas_con_error.xlsx',
                'content_base64': base64.b64encode(wb_data).decode('ascii'),
            }

        return success_response(
            {
                'total_rows': total,
                'created': created,
                'errors': errors,
                'valid_subjects': access_info['valid_subjects'],
                'denied_subjects': access_info['denied_subjects'],
                'partial_subject_access': access_info['partial_subject_access'],
                'errors_file': errors_file,
            },
            'Carga procesada.',
        )

    @extend_schema(
        summary='Descargar plantilla Excel',
        responses={200: OpenApiTypes.BINARY},
    )
    @action(detail=False, methods=['get'], url_path='template')
    def template(self, request):
        data = build_questions_template_workbook(request.user)
        resp = HttpResponse(
            data,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resp['Content-Disposition'] = 'attachment; filename="plantilla_preguntas.xlsx"'
        return resp
